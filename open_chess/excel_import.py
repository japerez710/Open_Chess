from __future__ import annotations

from pathlib import Path
from typing import Iterator


EXPECTED_HEADERS = ("alumno", "curso", "centro educativo")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalise_header(value: object) -> str:
    return " ".join(_text(value).casefold().split())


def _split_name(full_name: str, row_number: int) -> tuple[str, str]:
    if "," in full_name:
        last_name, first_name = (part.strip() for part in full_name.split(",", 1))
    else:
        parts = full_name.split(maxsplit=1)
        if len(parts) != 2:
            raise ValueError(
                f"Fila {row_number}: escribe el nombre y los apellidos del alumno."
            )
        first_name, last_name = parts
    if not first_name or not last_name:
        raise ValueError(
            f"Fila {row_number}: escribe el nombre y los apellidos del alumno."
        )
    return first_name, last_name


def _xlsx_rows(path: Path) -> Iterator[tuple[object, object, object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        yield from sheet.iter_rows(min_col=1, max_col=3, values_only=True)
    finally:
        workbook.close()


def _xls_rows(path: Path) -> Iterator[tuple[object, object, object]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(0)
        for row_index in range(sheet.nrows):
            values = [sheet.cell_value(row_index, column) if column < sheet.ncols else ""
                      for column in range(3)]
            yield values[0], values[1], values[2]
    finally:
        workbook.release_resources()


def read_student_rows(path: str | Path) -> list[dict[str, str]]:
    """Lee las tres primeras columnas de la primera hoja de un Excel."""
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".xlsx":
        rows = iter(_xlsx_rows(source))
    elif suffix == ".xls":
        rows = iter(_xls_rows(source))
    else:
        raise ValueError("Selecciona un archivo de Excel .xlsx o .xls.")

    try:
        headers = next(rows)
    except StopIteration as exc:
        raise ValueError("La hoja de Excel está vacía.") from exc
    normalised = tuple(_normalise_header(value) for value in headers)
    if normalised != EXPECTED_HEADERS:
        raise ValueError(
            "La fila 1 debe contener, en este orden: Alumno, Curso y Centro educativo."
        )

    students: list[dict[str, str]] = []
    errors: list[str] = []
    for row_number, values in enumerate(rows, start=2):
        full_name, course, center = (_text(value) for value in values)
        if not full_name and not course and not center:
            continue
        if not full_name:
            errors.append(f"Fila {row_number}: falta el nombre del alumno.")
            continue
        try:
            first_name, last_name = _split_name(full_name, row_number)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        students.append({
            "first_name": first_name,
            "last_name": last_name,
            "course": course,
            "center": center,
        })

    if errors:
        preview = "\n".join(errors[:8])
        extra = f"\n… y {len(errors) - 8} errores más." if len(errors) > 8 else ""
        raise ValueError(f"Corrige estos datos antes de importar:\n\n{preview}{extra}")
    if not students:
        raise ValueError("No hay alumnos para importar a partir de la fila 2.")
    return students
